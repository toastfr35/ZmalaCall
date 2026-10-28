"""Maintain release download instructions and the two phone Word authoring sources."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from release.guide_copy import PHONE, improve_readability

ROOT = Path(__file__).resolve().parents[1]
RELEASES = 'https://github.com/toastfr35/ZmalaCall/releases/latest'



def update_existing():
    replacements = {
        'terminal_prep.docx': 'ZmalaTerminal-Installation.pdf',
        'terminal_config.docx': 'ZmalaTerminal-Guide.pdf',
        'terminal_use.docx': 'ZmalaTerminal-Guide.pdf',
        'server_prep.docx': 'ZmalaServer-Installation.pdf',
    }
    for filename in ('terminal_prep.docx', 'terminal_config.docx', 'server_prep.docx'):
        path = ROOT / 'docs' / filename
        doc = Document(path)
        for p in doc.paragraphs:
            text = p.text
            for old, new in replacements.items():
                text = text.replace(old, new)
            if text.startswith('Open the newest successful Build terminal installer run'):
                text = ('Open the release link supplied by your administrator. Download '
                        'ZmalaTerminal-Pi4-setup.exe and the two ZmalaTerminal PDF guides. '
                        'Run the executable on your Windows PC; it installs the software on the Pi over Ethernet.')
            if text.startswith('Open the latest successful Build server installer run'):
                text = ('Open the release link supplied by your administrator. Download '
                        'ZmalaServer-Pi4-setup.exe and the two ZmalaServer PDF guides. '
                        'Run the executable on your Windows PC; the PDF guides sit beside it.')
            if text.startswith('A published release is another option'):
                text = ('Development builds are also available from successful GitHub Actions runs. '
                        'They may require GitHub sign-in and are separate from published releases.')
            if text != p.text:
                p.text = text
            # Replace old Actions download links without destroying hyperlink markup.
            for relation in list(p.part.rels.values()):
                if relation.is_external and '/actions/workflows/build-' in str(relation.target_ref):
                    relation._target = RELEASES
            for node in p._p.iter(qn('w:t')):
                if node.text and node.text.startswith('Latest successful '):
                    node.text = 'Published ZmalaCall releases'
        doc.core_properties.author = 'ZmalaCall'
        doc.core_properties.last_modified_by = ''
        doc.save(path)


def build_phone(filenames=None):
    for filename, (title, intro, sections) in PHONE.items():
        if filenames is not None and filename not in filenames:
            continue
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.8)
        for name, size in [('Normal', 11), ('Title', 27), ('Heading 1', 16)]:
            style = doc.styles[name]
            style.font.name = 'Calibri'
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string('1F2937')
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.line_spacing = 1.1
        for border in doc.styles.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('000000')
        doc.styles['Title'].font.color.rgb = RGBColor.from_string('000000')
        doc.styles['Heading 1'].paragraph_format.keep_with_next = True
        doc.add_paragraph(title, 'Title')
        doc.add_paragraph(intro)
        for heading, paragraphs in sections:
            paragraph = doc.add_heading(heading, 1)
            if heading in ('3 Allow incoming alerts', 'Missed calls and callbacks', 'Get help'):
                paragraph.paragraph_format.page_break_before = True
            for text in paragraphs:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.keep_together = True
                before, found, after = text.partition(RELEASES)
                paragraph.add_run(before)
                if found:
                    # Preserve a real link in Word and the derived release PDF.
                    link = OxmlElement('w:hyperlink')
                    link.set(qn('r:id'), paragraph.part.relate_to(
                        RELEASES, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                    run = OxmlElement('w:r')
                    value = OxmlElement('w:t')
                    value.text = found
                    run.append(value)
                    link.append(run)
                    paragraph._p.append(link)
                    paragraph.add_run(after)
        footer = section.footer.paragraphs[0]
        footer.add_run('ZmalaPhone  |  ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
        doc.core_properties.title = title
        doc.core_properties.author = 'ZmalaCall'
        doc.core_properties.last_modified_by = ''
        doc.save(ROOT / 'docs' / filename)


if __name__ == '__main__':
    update_existing()
    build_phone()
    improve_readability(ROOT)
