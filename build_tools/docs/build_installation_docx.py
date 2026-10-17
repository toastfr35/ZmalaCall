from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).resolve().parents[2]
# Extract just the connection diagram from the generated PDF, never whole pages.
# Keep all document text and the settings table editable in Word.
diagram=root/'tmp/pdfs/connections.png'
diagram.parent.mkdir(parents=True,exist_ok=True)
subprocess.run(['pdftoppm','-f','2','-l','2','-r','180','-x','110','-y','407',
                '-W','1268','-H','908','-singlefile','-png',
                str(root/'output/pdf/ZmalaCall-Pi4-Installation-Guide.pdf'),
                str(diagram.with_suffix(''))],check=True)
d=Document()
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
s=d.sections[0]
s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.65)
normal=d.styles['Normal'];normal.font.name='Segoe UI';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.1
for name,size in [('Title',24),('Heading 1',21),('Heading 2',13)]:
    st=d.styles[name];st.font.name='Segoe UI';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(9)
    st.font.bold=True
d.core_properties.title='ZmalaCall Raspberry Pi 4 installation guide'
d.core_properties.author='ZmalaCall'
footer=s.footer.paragraphs[0]
footer.add_run('ZmalaCall  |  Raspberry Pi 4  |  17 October 2026').font.size=Pt(8)
footer.add_run(' '*8+'Page ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def rich(label,body):
    a=d.add_paragraph();a.add_run(label+(' ' if label.endswith('.') else ': ')).bold=True;a.add_run(body);return a
def link(label,url):
    a=d.add_paragraph();el=OxmlElement('w:hyperlink');el.set(qn('r:id'),a.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1475BB');pr.append(co);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);el.append(r);a._p.append(el)
def pic(path,width,alt):
    a=d.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=a.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',alt)
def page(title):d.add_page_break();d.add_heading(title,1)

d.add_paragraph('ZmalaCall Raspberry Pi 4 installation guide','Title')
p('Install a new terminal from a freshly flashed SD card using a Windows PC. This guide covers the Argon ONE Pi 4 V2 case, the Nuroum C20 and the Microsoft Presenter+ remote.')
h('1 Prepare the SD card')
p('Use a Windows PC and an SD card reader. No keyboard or mouse is needed on the Pi.')
h('Get the software')
link('Download Raspberry Pi Imager from the official website','https://www.raspberrypi.com/software/')
p('Install the current Imager release. Select the OS from its catalogue so the correct customisation options are available.')
link('Download Zmala terminal setup from GitHub Actions','https://github.com/toastfr35/ZmalaCall/actions/workflows/build-terminal-installer.yml')
p('Sign in to GitHub. Open the latest successful main build, download the zmala-terminal-setup artifact and extract the ZIP. Run zmala_terminal_setup.exe.')
h('Choose the device OS and storage')
rich('Device','Raspberry Pi 4')
rich('OS','Raspberry Pi OS Lite (64-bit), Debian 13 Trixie')
rich('Storage','The microSD card you intend to erase')
h('Set these options before writing')
rows=[('Setting','Value'),('Hostname','zmala'),('Username','zmala'),('Password','zmala'),('SSH / remote access','Enabled, password authentication'),('Wi-Fi','Leave credentials unset; configure by QR later'),('Localisation','Your country, time zone and keyboard layout')]
table=d.add_table(rows=0,cols=2);table.autofit=False
table.columns[0].width=Inches(2);table.columns[1].width=Inches(4.9)
for i,row in enumerate(rows):
    cells=table.add_row().cells
    for cell,value in zip(cells,row):
        cell.text=value;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr=cell._tc.get_or_add_tcPr()
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'13263C' if i==0 else ('EEF4F8' if i%2 else 'FFFFFF'));pr.append(shade)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        pr.append(borders)
        for para in cell.paragraphs:
            para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
            for run in para.runs:
                run.font.size=Pt(10)
                if i==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    if i==0:table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
a=p('Choose the correct Wi-Fi country if Imager offers that field. Screen labels can vary between Imager versions.')
a.paragraph_format.space_before=Pt(6)
rich('Writing erases the selected card.','Double-check the drive, write the image and let verification finish. Safely eject the card, then insert it into the powered-off Pi.')
link('Official Imager and first boot instructions','https://www.raspberrypi.com/documentation/computers/getting-started.html')

page('2 Connect the equipment')
p('Argon ONE Pi 4 V2 case rear view. Connect power last.')
pic(root/'tmp/pdfs/connections.png',6.8,'Rear view of Argon ONE Pi 4 V2 showing HDMI to TV, USB-C power, remote charger in black USB 2, C20 in blue USB 3 and Ethernet to router.')
h('Before you switch on')
rich('Camera','Connect the C20 to a blue USB 3 socket.')
rich('Charger','Connect the Presenter cradle to the adjacent black USB socket.')
rich('TV','Connect one full-size HDMI socket on the case to the TV, then select that TV input. The second HDMI socket is unused.')
rich('Network','Keep Ethernet connected throughout installation. The Windows PC must be able to reach the Pi on the same local network.')
h('Now connect USB C power')
p('If the Pi does not start, press the case power button once. Allow the first boot to finish before checking it from Windows.')

page('3 Install ZmalaCall')
p('The Windows installer performs the transfer, installation, reboot and startup checks.')
h('Open the installer')
rich('Run zmala_terminal_setup.exe','No separate Python, SSH client or manual Pi commands are required. Keep the Pi powered and Ethernet connected until installation completes.')
h('Check the Pi')
p('The Pi hostname is zmala; its network address is zmala.local. Leave that address in the installer and click Check Pi. If it cannot be found, use its Ethernet IP from your router. The installer checks the connection, Pi model, OS, storage and wired connection.')
h('Confirm installation')
p('Argon ONE Pi 4 V2 fan setup is included automatically. Click Install and confirm the target in the dialog. Wait for the automatic reboot and final success message.')
h('Continue on the TV')
p('The Zmala boot image and first-start instructions should appear. A short blank screen during boot or HDMI negotiation can be normal.')
h('Already handled by the installer')
for t in ['Automatic startup and the Zmala boot screen.','Argon fan control during connected calls.','SSH restricted to Ethernet.','Reduced SD writes using RAM logs, RAM-only swap and temporary runtime files.']:p(t,'List Bullet')
h('If an error appears')
p('Keep power and Ethernet connected. Choose View details, then Open saved log. A lost connection does not necessarily stop installation on the Pi.')

page('4 Finish setup on the TV')
p('Have the Wi-Fi QR code and the terminal\'s server QR code ready. Keep the server running.')
h('Pair the Microsoft Presenter remote')
p('Switch the Presenter+ remote on. Turn it over and hold the rear Bluetooth button for about 3 seconds, until the white light flashes. Pairing happens automatically.')
pic(root/'zmala_terminal/assets/presenter-pairing.png',5.6,'Presenter+ front at left with green and red circular stickers covering the original symbols; back at right with the Bluetooth pairing button highlighted.')
p('After pairing, press and release Left, Right, Green, Red and Middle as prompted. QR setup starts only after all five buttons pass the test.')
h('Scan the two QR codes')
rich('1 Wi-Fi QR','Hold it in front of the C20 and wait for the connection to succeed.')
rich('2 Server QR','Scan it next to enroll the terminal. Its ID comes from this QR, not from the Pi hostname or the installer.')
rich('Replacing term2','Keep the old terminal stopped. Reuse its existing server QR if its address and credentials are still valid. Never run both terminals with the same identity.')
h('Final check')
p('Confirm that contacts appear and the remote buttons work. Make a test call and check picture and sound in both directions. Setup screens are in English; enrollment applies the selected language.')
p('If setup stalls, check the C20 cable, charge the remote, or check the scanned QR. The USB help screen appears automatically when the C20 is missing.')
out=root/'output/docx/ZmalaCall-Pi4-Installation-Guide.docx';out.parent.mkdir(parents=True,exist_ok=True);d.save(out);print(out)
