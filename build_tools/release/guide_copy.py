"""Reader-facing guide copy; applied by update-release-guides.py before PDF builds.

Keep UI labels literal. These instructions explain existing behavior, while the
normative contracts and technical detail remain in requirements/ and docs/*.md.
"""

PHONE = {
    'phone_installation.docx': (
        'Install ZmalaPhone',
        'Use this guide to install the Android app and connect your account. Ask the person '
        'setting up ZmalaCall for your private setup QR code. You do not need an administrator password.',
        [
            ('Before you start', [
                'You need an Android 9 or newer phone with a supported 64-bit processor, an internet '
                'connection and your own App setup QR. A Terminal QR is for a TV terminal and will not work here.',
                'Keep your QR private. Print it or show it on another screen for scanning. If you have a '
                'QR picture saved on this phone, you can choose Open QR image instead. If you received a '
                'PDF, display it on another screen or save a clear picture of the whole QR code.',
            ]),
            ('1 Download the app', [
                'Tap Download ZmalaPhone.apk in your setup document. Alternatively, open '
                'https://github.com/toastfr35/ZmalaCall/releases/latest and select ZmalaPhone.apk '
                'under Assets. You do not need a GitHub account.',
                'Open the downloaded file from Downloads or your file manager. If Android asks, allow '
                'that browser or file manager to install apps. Return to the file and tap Install. '
                'You can turn off this permission afterwards. This method needs neither a USB cable nor a SIM card.',
                'Already using a release version? Install over the existing app to keep your setup. '
                'Do not uninstall it for an ordinary update. If Android says the package conflicts, '
                'ask your helper before removing anything; a development copy may need different steps.',
            ]),
            ('2 Connect your account', [
                'Open ZmalaPhone. Tap Scan setup QR and point the camera at your code, or tap Open QR image '
                'and select a saved picture. Check the server address, then tap Connect.',
                'Allow camera and microphone access when asked. On Android 17, also allow local network '
                'access. A local HTTP address shows an unencrypted-connection warning; use it only on '
                'a trusted home network. For calls between different networks, your helper must supply an HTTPS setup code.',
                'Your contact list appears when setup succeeds. If the code is rejected, check that it is '
                'your current App QR and that the phone is connected to the internet or the required home network. '
                'Ask the person who supplied the code if it still fails.',
            ]),
            ('3 Allow incoming alerts', [
                'Read the call and callback notification explanation. Tap Continue, then Allow when Android asks. '
                'The bottom of the contact list shows whether incoming alerts are enabled.',
                'If it shows Enable notifications, tap it and allow notifications in Android settings. '
                'If it shows Try again, check your connection and tap it. Details provides help if the '
                'server is not ready to send alerts; ask your helper to check it.',
                'Open the app again after using Android Force stop, and at least once a month to renew '
                'incoming alerts. Battery-saving settings can delay alerts on some phones.',
            ]),
            ('4 Try a call together', [
                'Tap a contact card, then tap Call. Check that both people can hear each other and see '
                'moving video. Tap the video to show the controls, then tap End to hang up.',
                'Ask the other person to call you once while the app is open and once while the phone '
                'is on its Home screen or asleep. Answer each call and check sound and video. If alerts '
                'do not arrive, follow Receive calls in ZmalaPhone-Guide.pdf before relying on them.',
            ]),
        ],
    ),
    'phone_guide.docx': (
        'Use ZmalaPhone',
        'Tap someone’s contact card, then Call. This guide explains everyday calls, callbacks and '
        'updates. If the app still asks for a setup QR, follow ZmalaPhone-Installation.pdf first.',
        [
            ('Choose someone to call', [
                'The contact list stays upright in portrait. Tap a person’s card to open their page, '
                'then tap Call. Opening the card does not start a call. Swipe the list to see more people.',
                'Green Available means the person’s device was recently active. Blue Try calling means '
                'a phone alert may reach them, but this is not certain. You can try either. '
                'Red Offline means Call is unavailable; you can still open the card and request a callback.',
                'Tap Contacts or use Android Back to return to the list. The app follows the phone’s '
                'English or French language and supports Android’s larger text settings.',
            ]),
            ('Receive calls', [
                'An incoming call shows the caller’s name and picture. Tap Answer video call to accept '
                'or Decline to refuse. When the app is closed or the phone is asleep, use the incoming-call '
                'notification’s Answer or Video action.',
                'After answering, you can return Home, switch apps or lock the screen: sound continues '
                'and your camera pauses. Tap the ongoing-call notification to return. Your camera resumes '
                'when the call is visible again. The notification also lets you end the call.',
            ]),
            ('During a call', [
                'The large picture shows the other person; the small picture in the bottom-right corner '
                'shows you. Rotate the phone for a wider view. Tap the video to show the four controls. '
                'They usually hide again after five seconds.',
                'Mic: stop or restore the sound you send. Microphone muted stays visible while the other '
                'person cannot hear you.',
                'Sound: switch between loudspeaker, earpiece and muted listening. Phones without an '
                'earpiece skip that choice. Sound muted means you cannot hear the other person; it does '
                'not mute your microphone.',
                'Camera: switch between front and back cameras. End: finish the call.',
                'Swipe up on the video for louder sound or down for quieter sound. The phone’s volume '
                'buttons also work during a call. A new call uses loudspeaker at the existing Android call '
                'volume, so raise it here if it is too quiet.',
                'Video can take a few seconds after answer. Connecting video appears for up to fifteen '
                'seconds; a later message does not end the call. You can keep talking while waiting.',
            ]),
            ('Missed calls and callbacks', [
                'A missed-call count appears beside the contact. Open their page to call back or choose '
                'Clear missed calls. A connected call also clears that person’s count.',
                'To ask someone to call later, open their page and tap Request callback. You can do this '
                'for an offline contact, or accept the offer after an unanswered, declined or cancelled '
                'call. Wait for confirmation. The button then becomes Cancel request.',
                'Changed your mind? Tap Cancel request on the same person’s page. This cancels your '
                'pending request, not a request they sent you. If saving fails, check the connection and retry.',
                'Please call back means someone asked you to call them. Callback requested means you '
                'asked them. Incoming requests can also appear as Android notifications. Tapping one opens '
                'the contact without calling; choose Call when ready. A completed call, cancellation or '
                'expiry removes the reminder.',
            ]),
            ('Keep the app up to date', [
                'The version number is at the bottom of the contact list. Tap Updates, or Update available '
                'when shown. Choose Update to download, or Later to postpone for a day.',
                'Wait through Checking update, Preparing installation and Installing update. Allow '
                'installation from ZmalaPhone if Android asks, then confirm installation if required. '
                'Your setup is kept. Do not uninstall for an ordinary update.',
                'On supported Android 12 or newer phones, you can turn on Automatic updates on Wi-Fi. '
                'It starts off. It needs installation permission, an unmetered connection and thirty '
                'seconds without interaction on the visible contact list. Android may still ask for confirmation.',
                'Ringing and connected calls postpone updates. Leaving the app pauses a download; '
                'reopen Updates to continue. Installation briefly makes incoming calls unavailable. '
                'On Android 9 to 11, confirm updates manually.',
            ]),
            ('Get help', [
                'Cannot hear the other person: tap the video, check Sound is not muted and raise call '
                'volume. If they cannot hear you, check Mic is not muted and Android allows microphone access.',
                'No video: allow camera access, uncover the lens and wait briefly. If the picture does '
                'not arrive, end the call and try again. Tell your helper which device had no picture '
                'and whether sound worked.',
                'No incoming alerts: check internet access and the message at the bottom of Contacts. '
                'Use Enable notifications, Try again or Details as offered. Reopen after Android Force '
                'stop. Ask your helper to test calls with the screen locked and review battery-saving settings.',
                'Reconnecting: check the phone’s network and allow time to reconnect. A temporary server '
                'outage need not end an established call. Ask the person who supplied your QR about '
                'missing contacts or account access. Do not clear storage as a first troubleshooting step.',
            ]),
        ],
    ),
}


# Prefixes identify retained paragraphs, preserving diagrams, tables and links.
# Apply edits to the retained Word sources without disturbing their illustrations.
# The pass is idempotent so rebuilding never duplicates instructions.
REVISIONS = {
    'terminal_use.docx': {
        'Choose a person with the arrow buttons,': [
            'Use this guide for everyday calls after setup. Choose a familiar face with the arrow '
            'buttons, press green to confirm, and press red to finish.',
        ],
        'Sound and everyday help': [],
        'Press the middle button to show the call-volume overlay.': [
            'Press the middle button to show the volume bar. Press left for quieter sound or right '
            'for louder sound. Each click changes it by 5%. Press middle again to close the bar. '
            'You can do this during a call and keep seeing the other person.',
        ],
        'The call-volume overlay changes only': [
            'This changes the sound you hear during calls. It does not change your microphone, '
            'ringtone or button sounds. Ask your helper to adjust those separate sound settings if needed.',
        ],
        'Select an offline contact, press green and confirm': [
            'If a person is Offline, select their portrait and press green. Press green again '
            'to ask them to call you back, or red to cancel. A return-arrow badge shows your request.',
            'After an unanswered, declined or cancelled call, Zmala also offers a callback request. '
            'Press green to request it or red to return to your contacts.',
            'A large yellow Please call back flag means that person has asked you to call them. '
            'Choose them and call when ready. A request never starts a call automatically.',
        ],
        'To cancel your request,': [
            'To cancel a request you sent, select that person, press red, then confirm with green.',
        ],
        'A numbered badge can mean': [
            'A numbered badge marks missed calls or message notifications. Opening that person’s '
            'call or callback confirmation clears the count, even if you go back without calling.',
        ],
        'After ten seconds without a button press,': [
            'A thin blue bar at the bottom shows the countdown to returning to TV. A button press '
            'restarts the countdown, and calls pause it. Your helper chooses the delay. This works '
            'only when automatic TV control passed setup; otherwise use the TV remote.',
        ],
        'Webcam icon (if shown):': [
            'If a webcam symbol is shown, mint means power is on, a grey slash means off, and a '
            'question mark means checking. Pressing a button wakes sleeping equipment; allow a few '
            'seconds for the camera to reconnect.',
        ],
        'No picture or contact offline:': [
            'No picture during a call: wait a few seconds while video starts. If it stays missing, '
            'end the call and try again. Ask your helper if it keeps happening.',
            'Person unavailable: green Available means recently active; blue Try calling means '
            'a phone alert may reach them. Either can be called. Red Offline lets you request a callback instead.',
        ],
    },
    'terminal_prep.docx': {
        'Raspberry Pi 4 Model B in the Argon': [
            'Raspberry Pi 4 Model B in the Argon ONE Pi 4 V2 case, a 32GB or larger microSD card and a card reader.',
        ],
        'Development builds are also available': [],
        'Wait for the installer to report success.': [
            'Wait for the installer to report success. The TV should show the USB, remote or QR setup '
            'screen. Continue with the helper setup part of ZmalaTerminal-Guide.pdf. If the TV is blank, '
            'check its selected HDMI input and cable before repeating installation.',
        ],
    },
    'terminal_config.docx': {
        'For the person setting up the terminal': [
            'This part is for the helper setting up the TV terminal. Install the software first using '
            'ZmalaTerminal-Installation.pdf. Have the Presenter remote, TV remote and Wi-Fi and server '
            'QR documents ready. If an administrator has already supplied the codes, start at Pair the remote and connect Wi-Fi.',
        ],
        'Keep the QR documents private.': [
            'Keep both QR documents private: they grant access to the Wi-Fi and this terminal account. '
            'Use the server code made for this terminal. The Pi hostname does not select the account.',
        ],
        'Setup starts in English.': [
            'Setup starts in English and uses the chosen terminal language after the server code is '
            'accepted. Language choices are shown in their own language, in alphabetical order. '
            'Left and right wrap through the list. Some detailed help is in English.',
        ],
        'The terminal serial is shown at the bottom': [
            'The terminal ID appears at the bottom of Settings; give it to your helper if they need '
            'to find this device on the server. The C20 provides the camera, microphone and speaker automatically.',
            'If the C20 is waking, wait a few seconds. Zmala retries without switching sound to the '
            'TV. Press red if you want to cancel the call. Your administrator can adjust video quality '
            'and the starting call volume on the server.',
        ],
        'To choose a ringtone,': [
            'Open Ringtone with green. Press left or right to browse and hear each choice after a '
            'short pause. The list wraps at both ends. Press green to save or red to keep the previous tone.',
            'Ringtone volume controls incoming ringing, previews and the tone while waiting for someone to answer; 0% is silent. Button sounds volume is '
            'separate. Use middle and left/right to change the sound heard during calls. A call stops '
            'any ringtone preview.',
        ],
    },
    'server_prep.docx': {
        'For the person installing and administering': [
            'Use this guide to install and administer the shared Zmala server. You need a Windows PC '
            'and access to the router’s local network. The server Pi needs no screen, camera or remote. '
            'After installation, use ZmalaServer-Guide.pdf to add the family’s accounts and setup codes.',
        ],
        'The normal installation is for a trusted local network.': [
            'The normal installation is for a trusted local network. Leave router forwarding off for '
            'local-only calls. For calls from other networks, first complete the secure setup described '
            'below; forwarding the default unencrypted services is not sufficient.',
        ],
        'For an account portrait,': [
            'Choose a square portrait between 480 × 480 and 720 × 720 pixels. Use PNG or JPG for a '
            'still picture, up to 2 MiB, or GIF for animation, up to 16 MiB. The editor checks the picture '
            'and explains what to change if it is unsuitable.',
            'For GIFs, use 2 to 72 frames, a total length of at most five seconds and frame delays '
            'between 20 and 2000 ms. Zmala prepares a two-second animation at 20 fps. A terminal plays '
            'it once when the person is selected; a phone shows a still picture.',
        ],
        'In Terminals, choose native-language labels': [
            'Open Terminals, choose the device by its account name and change its preferences. '
            'Call volume, Ringtone volume and Button sounds volume are separate and use 5% steps. '
            'Remote & camera shows the remote’s battery and charger, and the C20’s connection and '
            'power when known. The Save changes bar stays visible as you scroll. Unsaved settings '
            'means your changes have not been applied; select Save changes to apply them.',
            'Open Setup to download its Wi-Fi and server QR documents. Give both privately to the '
            'person setting up that TV terminal, together with ZmalaTerminal-Guide.pdf.',
            'Open Callbacks to see Pending requests first. Choose History under Show callbacks '
            'to see completed and cancelled requests. The counts show how many are in '
            'each view. Cancel removes a pending request; looking at History changes nothing.',
        ],
        'To add ringtone choices,': [
            'In Terminals, open Ringtone. Preview plays on your computer. Add ringtones accepts audio '
            'you have permission to use, up to 12 MB per file, using the first 30 seconds. Choose the '
            'tone and Save changes. Nokia is always available; uploaded tones are included in backups.',
        ],
        'Connect your computer to the same local network': [
            'Open administration from the same local network as the server. It is deliberately blocked '
            'from other networks, even though phones and terminals can call remotely. After enabling '
            'secure mode, supply fresh HTTPS setup QR codes to every phone and terminal.',
            'Keep relay secrets and certificate keys in a separate private backup. The account backup '
            'does not include them.',
        ],
    },
}


# Also match retained guides already rewritten by an earlier generation.
REVISIONS['server_prep.docx']['Open Terminals, choose the device'] = (
    REVISIONS['server_prep.docx']['In Terminals, choose native-language labels'])

REVISIONS['server_prep.docx']['Open Callbacks to see Pending requests first.'] = [
    REVISIONS['server_prep.docx']['Open Terminals, choose the device'][2]]

def improve_readability(root):
    from docx import Document
    from docx.shared import Pt, RGBColor

    for filename, revisions in REVISIONS.items():
        path = root / 'docs' / filename
        doc = Document(path)
        for paragraph in list(doc.paragraphs):
            if paragraph._p.getparent() is None:
                continue  # A preceding replacement already removed this duplicate.
            for prefix, replacements in revisions.items():
                if not paragraph.text.startswith(prefix):
                    continue
                for existing in list(doc.paragraphs):
                    if existing._p is not paragraph._p and existing.text in replacements:
                        existing._p.getparent().remove(existing._p)
                for text in replacements:
                    new = paragraph.insert_paragraph_before(text, paragraph.style)
                    new.paragraph_format.keep_together = True
                paragraph._p.getparent().remove(paragraph._p)
                break
        if filename in ('server_prep.docx', 'terminal_config.docx'):
            # Keep administrator handover here; the phone guides own daily use.
            heading = next(p for p in doc.paragraphs if p.text == 'Setting up ZmalaPhone')
            following = heading._p.getnext()
            while following is not None and following.tag.endswith('}p'):
                next_element = following.getnext()
                following.getparent().remove(following)
                following = next_element
            for text in (
                'Create an App account for each phone and add it to the family’s shared group. '
                'Keep Terminal accounts for TV terminals. In Accounts, select the chain-link '
                'Connection details button beside the person, then Download server QR PDF. '
                'Give this private code to that phone user; they do not need the administration password.',
                'Send the signed ZmalaPhone.apk download link with ZmalaPhone-Installation.pdf '
                'and ZmalaPhone-Guide.pdf. The installation guide covers a printed QR or a QR '
                'picture saved on the same phone, permissions and the first test call.',
                'For incoming alerts while the app is closed, configure the server’s private Firebase '
                'sender and ask the phone user to allow notifications. Test a call with their app '
                'open and again with the phone asleep. For calls from other networks, configure '
                'secure relay mode first and supply HTTPS setup codes.',
            ):
                paragraph = doc.add_paragraph(text)
                paragraph.paragraph_format.keep_together = True
        if filename == 'server_prep.docx':
            if not any(p.text == 'Developer access to a terminal' for p in doc.paragraphs):
                doc.add_heading('Developer access to a terminal', level=1)
                for text in (
                    'For remote TV-plugin testing, open Terminals, choose Developer access on the '
                    'required terminal and select Enable for this terminal. Only one terminal can '
                    'have access at a time. The dialog shows whether its tunnel has connected.',
                    'Developer access stays enabled until you select Disable developer access. '
                    'Disabling disconnects existing sessions. Keep your development computer on '
                    'the same subnet as the server and use your normal server and Pi SSH logins.',
                    'For a terminal on another network, forward TCP port 2222 on the server router '
                    'to TCP 2222 on the server Pi. This separate port accepts only the tunnel '
                    'account. Never forward ordinary SSH port 22. Point the C20 at the TV so '
                    'the developer can check the displayed picture while testing.',
                ):
                    doc.add_paragraph(text).paragraph_format.keep_together = True
            ports = [
                ('Purpose', 'Protocol', 'Public port', 'Pi port'),
                ('HTTPS setup', 'TCP', '443', '8443'),
                ('Encrypted signalling', 'TCP', '5061', '5061'),
                ('Encrypted TURN relay', 'TCP', '5349', '5349'),
                ('Relay media', 'UDP', '30000-30199', '30000-30199'),
                ('Certificate renewal using HTTP-01', 'TCP', '80', '80'),
            ]
            if not any(t.cell(0, 0).text == 'Purpose' for t in doc.tables):
                anchor = next(p for p in doc.paragraphs if p.text.startswith('Find Port forwarding,'))
                table = doc.add_table(rows=0, cols=4)
                table.style = 'Table Grid'
                for values in ports:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = value
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                anchor._p.addnext(table._tbl)
                note = doc.add_paragraph(
                    'These ports are an example for the single-Pi secure setup. Match 5061 and 5349 '
                    'to the advertised SIP and TURN ports. Match the UDP range to coturn’s min-port '
                    'and max-port; it is needed even when callers use TURN over TLS. Port 80 is only '
                    'for the certificate challenge responder, not administration. Keep it available '
                    'for renewal if using HTTP-01; DNS validation does not need this rule.')
                table._tbl.addnext(note._p)
                for paragraph in list(doc.paragraphs):
                    if paragraph.text.startswith(('HTTPS configuration:', 'Encrypted call signalling:',
                            'Encrypted TURN relay:', 'Relay media:', 'Certificate validation, if using')):
                        paragraph._p.getparent().remove(paragraph._p)
        # Keep the familiar illustrations and table layout; make all headings
        # consistent and step labels distinguishable without relying on colour.
        release_link_seen = False
        for paragraph in doc.paragraphs:
            if paragraph.text == 'Published ZmalaCall releases':
                if release_link_seen:
                    paragraph._p.getparent().remove(paragraph._p)
                    continue
                release_link_seen = True
            if paragraph.text.startswith('These ports are an example'):
                paragraph.paragraph_format.keep_together = True
            if paragraph.text.startswith('No sound:'):
                paragraph.text = 'No sound: increase the volume and check that the C20 is connected.'
            if paragraph.text.startswith('Remote does not respond:') and 'lightning bolt' not in paragraph.text:
                paragraph.add_run(' The lightning bolt beside its battery means the charger is connected.')
            if paragraph.text[:1].isdigit() and '  ' in paragraph.text[:5]:
                paragraph.style = 'Heading 2'
            if paragraph.style.name.startswith(('Title', 'Heading')):
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
        for style in ('Title', 'Heading 1', 'Heading 2'):
            doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles['Normal'].font.size = Pt(11)
        doc.core_properties.author = 'ZmalaCall'
        doc.core_properties.last_modified_by = ''
        doc.save(path)
