"""Append the secure-network operator notes to the retained Word user guides.

Run with the document workspace runtime. Existing guide structure is preserved;
this section is replaced by title on subsequent runs.
"""
from pathlib import Path
import sys
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SECTIONS = {
 'server_prep.docx': [
  'Calls from networks that block UDP',
  'The normal installation serves the local network. For terminals in separate buildings, ask the network administrator to configure secure relay mode before printing the server QR codes.',
  'This mode needs public endpoints, trusted certificates and a TURN relay. For one home Internet connection, the usual public TCP ports are 443 for configuration, 5061 for call signalling and 5349 for the relay. See Router port forwarding for remote calls in ZmalaServer-Installation.pdf. The technical setup is described in docs/server.md; it is not enabled automatically by the installer.',
  'Connect your computer to the same local network as the server before opening the HTTPS administration address. Administration, including sign-in, is blocked from other networks. Phones and terminals can still connect remotely for calls. After enabling secure mode, print new server QR codes and enroll both terminals with them. Keep the private relay secret and certificate keys in a separate secure backup; the account backup does not include them.',
  'Make a call in each direction and check sound and moving video before handing over the devices. A network that allows only web browsing through a proxy can still block calls. Give its administrator the configured endpoint names and TCP ports. Using TCP 443 for all three services requires a suitable separate-endpoint deployment; a normal web proxy cannot carry SIP or TURN.'
 ],
 'terminal_config.docx': [
  'Using a secure server connection',
  'If the terminal will be used on a network that blocks UDP, ask the server administrator for an HTTPS server QR code. The administrator must first configure the secure call service and media relay.',
  'Install the current terminal software, then use Settings and Reset to enroll with the new QR code. An old HTTP QR code cannot activate secure relay mode. Do not turn off certificate checks to work around an enrollment error.',
  'Call and answer controls stay the same. The terminal receives temporary relay credentials from the server. It updates them while idle, so an active call is not interrupted. If the terminal has been offline for a long time, restore its server connection before making a call.',
  'Before use, make a call in both directions and check sound and video. If registration or media fails, ask the administrator to check the endpoint addresses, certificates, clock and network access. No automatic switch to an unencrypted connection is made.'
 ]
}
FORWARDING = [
 ('Heading 1', 'Router port forwarding for remote calls'),
 ('Normal', 'Use this section when phones or terminals will call from outside the server network. Local-only installations need no forwarding. This example uses one public IPv4 address, one hostname, and the Zmala server and TURN relay on the same Pi. Forwarding alone does not enable secure calls: arrange the secure service, relay and certificates described in docs/server.md first.'),
 ('Heading 2', 'Reserve the Pi address'),
 ('Normal', 'Open your router administration page from its local network. Find LAN, DHCP or Connected devices, select the Pi Ethernet connection and create an address reservation, sometimes called a static lease. For example, reserve 192.168.1.50 if it belongs to your router subnet and is available. Use your actual reserved address in every rule below. Restart the Pi if required and check that its address stays the same.'),
 ('Heading 2', 'Check the public connection'),
 ('Normal', 'Set a public DNS or dynamic DNS hostname to the router public IPv4 address. Configure a DDNS updater if the provider changes that address. Check the router WAN or Internet address: a private address, or an address in 100.64.0.0 to 100.127.255.255, can indicate another router or carrier-grade NAT. Forwarding on this router alone will then be insufficient.'),
 ('Normal', 'With two routers you control, use bridge mode on the upstream router or forward the same public ports through it to the downstream router, then to the Pi. With provider carrier-grade NAT, ask for a publicly reachable IPv4 address or use a separately planned public relay deployment.'),
 ('Heading 2', 'Add the forwarding rules'),
 ('Normal', 'Find Port forwarding, NAT, Virtual server or Port mapping. Add one rule per entry below. Select the stated protocol, enter the public port as External or WAN port, set the destination to the reserved Pi address, and enter the destination port as Internal or LAN port. Enable and save each rule. The remote source may need to allow any address for callers on changing networks; the destination is always the Pi.'),
 ('Normal', 'HTTPS configuration: TCP external 443 to internal 8443. Remote apps and terminals use this for enrollment and configuration. Administration on the same HTTPS service must remain restricted to the server LAN.'),
 ('Normal', 'Encrypted call signalling: TCP external 5061 to internal 5061. The server secure policy must advertise public SIP port 5061.'),
 ('Normal', 'Encrypted TURN relay: TCP external 5349 to internal 5349. Configure coturn to listen on TLS port 5349 and the server policy to advertise public TURN port 5349.'),
 ('Normal', 'Relay media: UDP external 30000-30199 to the same internal range 30000-30199. Replace both ranges if coturn uses different min-port and max-port values. This UDP range is still needed when clients connect to TURN over TCP/TLS.'),
 ('Normal', 'Certificate validation, if using HTTP-01: TCP external 80 to internal 80, where the ACME challenge responder runs. Keep the route available for automatic renewal; the responder may run only during validation. DNS-based validation does not require this rule. Do not send port 80 to the unencrypted administration service.'),
 ('Heading 2', 'Match the server configuration'),
 ('Normal', 'Match the rules to the actual services and Pi firewall. Configure coturn with the public-to-private address mapping and the same relay range; keep the mapping current when the public address changes. Configure trusted hostname certificates and automatic renewal. Keep administration LAN-only. Do not forward SSH port 22, plain HTTP administration port 8080, plain SIP port 5060 or the router administration page, and do not place the Pi in a router DMZ.'),
 ('Heading 2', 'Check from another network'),
 ('Normal', 'Give devices fresh HTTPS setup QR codes using the public hostname. Use a phone on mobile data with Wi-Fi off, or another outside network, and call in each direction. Check moving video and speech on both devices, then an incoming call with the app closed and notifications enabled. No forwarding is needed on the phone or terminal router.'),
 ('Normal', 'Opening the public administration page from outside the LAN should return HTTP 403; this is expected and does not mean device APIs are blocked. Inside the LAN, check administration using the same HTTPS hostname. If the router cannot loop the public address back into the LAN, arrange local DNS to the Pi and a matching local HTTPS listener; DNS alone does not translate port 443 to 8443. Do not disable certificate checks.'),
 ('Normal', 'If calls fail, check the Pi reservation, DNS address, rule protocols and ranges, server listeners, Pi firewall and certificate expiry. A port checker only tests an active listener and cannot prove that video works. Keep a private record of the rules and test again after router replacement or a public address change.'),
]


def update_forwarding(doc):
 """Replace the operator section without changing surrounding content."""
 removing = False
 for paragraph in list(doc.paragraphs):
  if paragraph.text == FORWARDING[0][1]:
   removing = True
  elif removing and paragraph.style.name == 'Heading 1':
   removing = False
  if removing:
   paragraph._element.getparent().remove(paragraph._element)
 anchor = next(p for p in doc.paragraphs if p.text == 'Create contacts and verify the installation')
 for style, text in FORWARDING:
  paragraph = anchor.insert_paragraph_before(text, style)
  if style == 'Heading 1':
   paragraph.paragraph_format.page_break_before = True


for filename, paragraphs in SECTIONS.items():
 if len(sys.argv) > 1 and filename not in sys.argv[1:]:
  continue
 path = ROOT/'docs'/filename
 doc = Document(path)
 removing = False
 for paragraph in list(doc.paragraphs):
  if paragraph.text == paragraphs[0]:
   removing = True
  elif removing and paragraph.style.name == 'Heading 1':
   removing = False
  if removing:
   paragraph._element.getparent().remove(paragraph._element)
 heading = doc.add_paragraph(paragraphs[0], style='Heading 1')
 heading.paragraph_format.page_break_before = True
 for text in paragraphs[1:]:
  doc.add_paragraph(text)
 if filename == 'server_prep.docx':
  update_forwarding(doc)
 doc.save(path)
