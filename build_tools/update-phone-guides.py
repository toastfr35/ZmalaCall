"""Update retained operator guides for App/Terminal enrollment (document workspace runtime)."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TITLE = 'Setting up ZmalaPhone'
SECTIONS = {'server_prep.docx': ['Create an App account for each Android phone, separate from Pi Terminal accounts, and add '
                      'it to a shared calling group. Give the account’s private setup QR PDF to the phone user by '
                      'email or on paper; they do not need server administration access.',
                      'Download the signed ZmalaPhone.apk from the supplied release link. Open it and allow '
                      'installation when Android asks. In ZmalaPhone, choose Scan setup QR for a printed code or '
                      'another screen, or Open QR image for a saved picture on the same phone. Confirm the server '
                      'address. Allow camera, microphone and, on Android 17, local network access.',
                      'Allow notifications for incoming calls while the app is closed or the phone sleeps. The '
                      'footer offers notification settings, retry or help when alerts are unavailable. Alerts '
                      'need internet and the server’s private Firebase sender configuration. Reopen after Android '
                      'Force stop and at least monthly.',
                      'Each contact is one card opening that person’s Call, Request callback or Cancel request '
                      'actions and missed-call information. Green Available means recently active; blue Try '
                      'calling means push may work but is unconfirmed. Red Offline cards still open, with Call '
                      'disabled. The app supports English/French and larger Android text.',
                      'During calls, tap video for labeled microphone, listening, camera-switch and hangup '
                      'controls. Swipe up/down or use volume keys for listening volume. Home and screen locking '
                      'retain audio and pause camera sending. The ongoing notification returns to or ends the '
                      'call. Check real calls in both directions before handover.',
                      'See ZmalaPhone-Installation.pdf and ZmalaPhone-Guide.pdf for complete instructions. Signed '
                      'updates keep setup. To replace setup, use Android Clear storage and a fresh App QR; the '
                      'server account remains unchanged.',
                      'For different-network calls, first configure secure relay mode, certificates and a stable '
                      'public hostname as described in docs/server.md. Supply HTTPS setup QR codes. Keep '
                      'administration LAN-only and keep server keys private.'],
 'terminal_config.docx': ['Each Pi terminal needs a Terminal account and each phone its own App account in a '
                          'shared calling group. Keep the Pi account as Terminal when adding a phone. Ask the '
                          'administrator for the signed APK link and private App setup QR; phone users do not '
                          'need server administration access.',
                          'Download the signed ZmalaPhone.apk from the supplied release link. Open it and allow '
                          'installation when Android asks. In ZmalaPhone, choose Scan setup QR for a printed code '
                          'or another screen, or Open QR image for a saved picture on the same phone. Confirm the '
                          'server address. Allow camera, microphone and, on Android 17, local network access.',
                          'Allow notifications for incoming calls while the app is closed or the phone sleeps. '
                          'The footer offers notification settings, retry or help when alerts are unavailable. '
                          'Alerts need internet and the server’s private Firebase sender configuration. Reopen '
                          'after Android Force stop and at least monthly.',
                          'Each contact is one card opening that person’s Call, Request callback or Cancel '
                          'request actions and missed-call information. Green Available means recently active; '
                          'blue Try calling means push may work but is unconfirmed. Red Offline cards still open, '
                          'with Call disabled. The app supports English/French and larger Android text.',
                          'During calls, tap video for labeled microphone, listening, camera-switch and hangup '
                          'controls. Swipe up/down or use volume keys for listening volume. Home and screen '
                          'locking retain audio and pause camera sending. The ongoing notification returns to or '
                          'ends the call. Check real calls in both directions before handover.',
                          'See ZmalaPhone-Installation.pdf and ZmalaPhone-Guide.pdf for complete instructions. '
                          'Signed updates keep setup. To replace setup, use Android Clear storage and a fresh App '
                          'QR; the server account remains unchanged.']}

def update():
    for filename, paragraphs in SECTIONS.items():
        path = ROOT / 'docs' / filename
        doc = Document(path)
        removing = False
        for paragraph in list(doc.paragraphs):
            if paragraph.text == TITLE:
                removing = True
            elif removing and paragraph.style.name == 'Heading 1':
                removing = False
            if removing:
                paragraph._element.getparent().remove(paragraph._element)
            elif paragraph.text.startswith('In Accounts, create an account for each terminal or phone,'):
                paragraph.text = 'In Accounts, create a Terminal account for each Pi terminal and an App account for each Android phone. Set its display name and portrait, keep generated credentials private and supply real email addresses where reminders are needed.'
        heading = doc.add_paragraph(TITLE, style='Heading 1')
        heading.paragraph_format.page_break_before = True
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(path)

if __name__ == '__main__':
    update()
